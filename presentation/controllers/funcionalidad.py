import os
import platform
import subprocess
import threading
import time
import wave
import contextlib

import tkinter as tk
from tkinter import filedialog

from presentation.views.app_dialog import ask_yes_no_cancel, show_error, show_info, show_warning

from mutagen import File

import config
from config import (
    logger,
    report_error,
    idiomas,
    ffmpeg_path,
)


def _ask_on_main_thread(ventana, title, message):
    """Show a yes/no/cancel dialog on the main thread.

    Blocks the calling (worker) thread until the user responds,
    without freezing the Tkinter event loop.

    Returns True (Yes), False (No), or None (Cancel).
    """
    result = [None]
    event = threading.Event()

    def _show():
        result[0] = ask_yes_no_cancel(ventana, title, message)
        event.set()

    ventana.after(0, _show)
    event.wait()
    return result[0]


def _show_error_on_main_thread(ventana, title, message):
    """Show an error dialog on the main thread (non-blocking for worker)."""
    ventana.after(0, lambda: show_error(ventana, title, message))


def _show_info_on_main_thread(ventana, title, message):
    """Show an info dialog on the main thread (non-blocking for worker)."""
    ventana.after(0, lambda: show_info(ventana, title, message))


def convertir_a_wav(audio_path):
    """Convert an audio file to WAV with the bundled ffmpeg (local, offline).

    Used only as a playback fallback when pygame cannot decode a format
    directly. Transcription does not need this: faster-whisper decodes files
    on its own.
    """
    try:
        logger.info(f"Intentando convertir: {audio_path}")
        base, _ = os.path.splitext(audio_path)
        output_path = base + ".wav"

        if os.path.exists(output_path):
            logger.info(f"El archivo WAV ya existe: {output_path}")
            return output_path

        command = [ffmpeg_path, "-i", audio_path, output_path]

        startupinfo = subprocess.STARTUPINFO()
        if platform.system() == "Windows":
            startupinfo.dwFlags |= subprocess.STARTF_USESHOWWINDOW
            startupinfo.wShowWindow = subprocess.SW_HIDE

        subprocess.run(
            command,
            startupinfo=startupinfo,
            check=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
        )

        logger.info(f"Archivo convertido a WAV: {output_path}")
        return output_path
    except Exception as e:
        logger.error(f"Error al convertir archivo a WAV: {str(e)}")
        raise


def obtener_duracion_audio(ruta_archivo):
    try:
        audio = File(ruta_archivo)
        if audio is not None:
            return int(audio.info.length)
    except Exception:
        pass

    if ruta_archivo.lower().endswith(".wav"):
        try:
            with contextlib.closing(wave.open(ruta_archivo, "r")) as f:
                frames = f.getnframes()
                rate = f.getframerate()
                duration = frames / float(rate)
                return int(duration)
        except Exception:
            pass

    return 0


def seleccionar_archivos(lista_archivos, lista_archivos_paths):
    if getattr(seleccionar_archivos, "_lock", False):
        logger.warning("seleccionar_archivos está bloqueado para evitar apertura doble.")
        return
    setattr(seleccionar_archivos, "_lock", True)
    
    try:
        file_paths = filedialog.askopenfilenames(
            filetypes=[
                ("Archivos de Audio", "*.mp3 *.wav *.flac *.ogg *.m4a *.mp4 *.aac *.opus")
            ],
            title="Seleccionar archivos de audio",
        )
        logger.info(f"Archivos seleccionados por el diálogo: {file_paths} (tipo: {type(file_paths)})")
        
        if not file_paths:
            return

        # Si por alguna razón Tcl/Tk devuelve un string en vez de una lista/tupla,
        # lo convertimos de forma segura usando shlex.
        if isinstance(file_paths, str):
            import shlex
            try:
                file_paths = shlex.split(file_paths)
            except Exception:
                file_paths = [file_paths]

        archivos_no_agregados = []
        for file_path in file_paths:
            # Limpiar llaves de Tcl {} que a veces envuelven rutas con espacios en Windows
            file_path = file_path.strip("{}").strip()
            if not file_path:
                continue
                
            file_name = os.path.basename(file_path)
            duracion = obtener_duracion_audio(file_path)
            duracion_str = format_duration(duracion)
            item = f"{file_name} ({duracion_str})"
            
            logger.info(f"Procesando archivo para lista: path={file_path}, name={file_name}, duration={duracion_str}")
            
            if item not in lista_archivos.get(0, tk.END):
                lista_archivos.insert(tk.END, item)
                lista_archivos_paths[file_path] = item
            else:
                archivos_no_agregados.append(file_name)
                
        if archivos_no_agregados:
            show_warning(
                lista_archivos.winfo_toplevel(),
                "Archivos Duplicados",
                "Los siguientes archivos ya estaban en la lista y no se añadieron nuevamente:\n"
                f"{', '.join(archivos_no_agregados)}",
            )
    finally:
        # Liberar el bloqueo después de 300ms para ignorar eventos duplicados
        lista_archivos.after(300, lambda: setattr(seleccionar_archivos, "_lock", False))


def limpiar(text_area):
    text_area.delete("1.0", tk.END)


def ajustar_texto_sencillo(texto, max_ancho=90):
    palabras = texto.split()
    lineas = []
    linea_actual = []
    longitud_actual = 0

    for palabra in palabras:
        if len(palabra) > max_ancho:
            if linea_actual:
                lineas.append(" ".join(linea_actual))
            for i in range(0, len(palabra), max_ancho):
                lineas.append(palabra[i: i + max_ancho])
            linea_actual = []
            longitud_actual = 0
        elif longitud_actual + len(palabra) + (1 if linea_actual else 0) <= max_ancho:
            linea_actual.append(palabra)
            longitud_actual += len(palabra) + (1 if linea_actual else 0)
        else:
            lineas.append(" ".join(linea_actual))
            linea_actual = [palabra]
            longitud_actual = len(palabra)

    if linea_actual:
        lineas.append(" ".join(linea_actual))

    if len(lineas) > 1 and len(lineas[-1]) < max_ancho // 2:
        penultima = lineas[-2].split()
        ultima = lineas[-1].split()
        while penultima and len(" ".join(penultima + [ultima[0]])) <= max_ancho:
            ultima.insert(0, penultima.pop())
        lineas[-2] = " ".join(penultima)
        lineas[-1] = " ".join(ultima)

    return "\n".join(lineas)


def contar_palabras_y_inaudibles(texto):
    palabras = texto.split()
    inaudibles = palabras.count("[inaudible]")
    palabras_sin_inaudibles = len(palabras) - inaudibles
    return palabras_sin_inaudibles, inaudibles


def parse_time_to_srt(time_str: str) -> str:
    parts = time_str.split(":")
    if len(parts) == 3:
        hours, mins, secs = int(parts[0]), int(parts[1]), int(parts[2])
    else:
        hours = 0
        mins, secs = int(parts[0]), int(parts[1])
    return f"{hours:02d}:{mins:02d}:{secs:02d},000"

def parse_time_to_vtt(time_str: str) -> str:
    parts = time_str.split(":")
    if len(parts) == 3:
        hours, mins, secs = int(parts[0]), int(parts[1]), int(parts[2])
    else:
        hours = 0
        mins, secs = int(parts[0]), int(parts[1])
    return f"{hours:02d}:{mins:02d}:{secs:02d}.000"

def text_to_srt(text: str) -> str:
    import re
    lines = text.strip().split("\n")
    srt_parts = []
    index = 1
    
    # Matches [MM:SS - MM:SS] or [HH:MM:SS - HH:MM:SS]
    pattern = re.compile(r"^\[((?:\d{1,2}:)?\d{1,2}:\d{2})\s*-\s*((?:\d{1,2}:)?\d{1,2}:\d{2})\]\s*(.*)$")
    
    for line in lines:
        match = pattern.match(line.strip())
        if match:
            start_str, end_str, content = match.groups()
            try:
                start_srt = parse_time_to_srt(start_str)
                end_srt = parse_time_to_srt(end_str)
                srt_parts.append(f"{index}\n{start_srt} --> {end_srt}\n{content}\n")
                index += 1
            except Exception:
                pass
    return "\n".join(srt_parts)

def text_to_vtt(text: str) -> str:
    import re
    lines = text.strip().split("\n")
    vtt_parts = ["WEBVTT\n"]
    
    pattern = re.compile(r"^\[((?:\d{1,2}:)?\d{1,2}:\d{2})\s*-\s*((?:\d{1,2}:)?\d{1,2}:\d{2})\]\s*(.*)$")
    
    for line in lines:
        match = pattern.match(line.strip())
        if match:
            start_str, end_str, content = match.groups()
            try:
                start_vtt = parse_time_to_vtt(start_str)
                end_vtt = parse_time_to_vtt(end_str)
                vtt_parts.append(f"\n{start_vtt} --> {end_vtt}\n{content}\n")
            except Exception:
                pass
    return "\n".join(vtt_parts)

def exportar_transcripcion(transcripcion_resultado):
    if not transcripcion_resultado or not transcripcion_resultado.strip():
        show_warning(None, "Advertencia", "No hay transcripción para exportar.")
        return
        
    output_file = filedialog.asksaveasfilename(
        defaultextension=".txt",
        filetypes=[
            ("Archivo de texto", "*.txt"),
            ("Documento de Word", "*.docx"),
            ("Subtítulos SRT", "*.srt"),
            ("Subtítulos VTT", "*.vtt")
        ],
        title="Guardar transcripción como",
    )
    
    if not output_file:
        return
        
    _, ext = os.path.splitext(output_file.lower())
    
    try:
        if ext == ".docx":
            import docx
            doc = docx.Document()
            doc.add_heading("Transcripción de AudioText", 0)
            for paragraph in transcripcion_resultado.split("\n"):
                if paragraph.strip():
                    doc.add_paragraph(paragraph)
            doc.save(output_file)
        elif ext == ".srt":
            srt_content = text_to_srt(transcripcion_resultado)
            if not srt_content.strip():
                show_warning(
                    None,
                    "Advertencia",
                    "No se detectaron marcas de tiempo [MM:SS - MM:SS] en el texto.\n\n"
                    "Se guardará como archivo de texto plano pero con extensión .srt.",
                )
                srt_content = transcripcion_resultado
            with open(output_file, "w", encoding="utf-8") as f:
                f.write(srt_content)
        elif ext == ".vtt":
            vtt_content = text_to_vtt(transcripcion_resultado)
            if len(vtt_content.strip()) <= 7: # only contains WEBVTT\n
                show_warning(
                    None,
                    "Advertencia",
                    "No se detectaron marcas de tiempo [MM:SS - MM:SS] en el texto.\n\n"
                    "Se guardará como archivo de texto plano pero con extensión .vtt.",
                )
                vtt_content = "WEBVTT\n\n" + transcripcion_resultado
            with open(output_file, "w", encoding="utf-8") as f:
                f.write(vtt_content)
        else: # .txt
            with open(output_file, "w", encoding="utf-8") as f:
                f.write(transcripcion_resultado)
                
        show_info(None, "Información", f"Transcripción guardada en {output_file}.")
        logger.info(f"Transcripción guardada en {output_file}.")

        try:
            if platform.system() == "Darwin":
                subprocess.call(("open", output_file))
            elif platform.system() == "Windows":
                os.startfile(output_file)
            else:
                subprocess.call(("xdg-open", output_file))
            logger.info(f"Archivo abierto: {output_file}")
        except Exception as e:
            logger.error(f"Error al abrir el archivo: {str(e)}")
            
    except Exception as e:
        show_error(None, "Error de exportación", f"No se pudo exportar el archivo:\n\n{e}")
        logger.error(f"Error al exportar archivo: {e}")



def borrar_archivo(lista_archivos, lista_archivos_paths):
    seleccion = lista_archivos.curselection()
    if seleccion:
        indice = seleccion[0]
        archivo = lista_archivos.get(indice)
        lista_archivos.delete(indice)
        lista_archivos_paths.pop(
            next(
                key for key, value in lista_archivos_paths.items() if value == archivo
            ),
            None,
        )
        return True
    return False


def format_duration(seconds: float) -> str:
    """Format seconds as MM:SS or H:MM:SS for durations >= 1 hour."""
    total = int(seconds)
    hours = total // 3600
    minutes = (total % 3600) // 60
    secs = total % 60
    if hours > 0:
        return f"{hours}:{minutes:02d}:{secs:02d}"
    else:
        return f"{minutes:02d}:{secs:02d}"


format_time = format_duration


def iniciar_transcripcion_thread(
    lista_archivos,
    text_area,
    archivo_procesando,
    lista_archivos_paths,
    transcripcion_resultado,
    progress_bar,
    ventana,
    boton_transcribir,
    combobox_idioma_entrada,
    combobox_idioma_salida,
    diarizar_var=None,
    marcas_tiempo_var=None,
    spinner=None,
    frame_progress=None,
    on_file_done=None,
):
    from infrastructure.audio.reproductor import reproductor

    if reproductor.reproduciendo:
        show_warning(
            ventana,
            "Advertencia",
            "Hay una reproduccion en curso. Por favor, detenga la reproduccion antes de transcribir.",
        )
        return

    seleccion = lista_archivos.curselection()
    if not seleccion:
        show_warning(ventana, "Advertencia", "Seleccione un archivo de audio para transcribir.")
        config.transcripcion_en_curso = False
        return

    if config.transcripcion_activa:
        config.transcripcion_activa = False
        config.transcripcion_en_curso = False
        boton_transcribir.configure(text="🎙️   Transcribir")
        progress_bar.configure(value=0)
        progress_bar.pack_forget()
        if spinner:
            spinner.stop()
            spinner.pack_forget()
        if frame_progress:
            frame_progress.pack_forget()
    else:
        # Read GUI variables safely on the main GUI thread before spawning the background thread
        idioma_entrada_val = combobox_idioma_entrada.get()
        idioma_salida_val = combobox_idioma_salida.get()
        diarizar_val = False
        timestamps_val = bool(marcas_tiempo_var.get()) if marcas_tiempo_var is not None else False
        
        # Safe check for diarization availability on main GUI thread
        if diarizar_val and not diarizer.is_available():
            show_info(
                ventana,
                "Diferenciar hablantes no disponible",
                "Falta whisperx o el token de HuggingFace (HF_TOKEN en .env). "
                "Se transcribirá sin diferenciar hablantes.",
            )
            diarizar_val = False

        config.transcripcion_activa = True
        config.transcripcion_en_curso = True
        boton_transcribir.configure(text="⏹   Detener")
        progress_bar["value"] = 0
        if frame_progress:
            frame_progress.pack(side=tk.TOP, fill=tk.X, pady=(10, 0))
        progress_bar.pack(pady=5, padx=60, fill=tk.X)
        if spinner:
            spinner.start()
            spinner.pack(side=tk.LEFT, padx=5)
        threading.Thread(
            target=iniciar_transcripcion,
            args=(
                lista_archivos,
                text_area,
                archivo_procesando,
                lista_archivos_paths,
                transcripcion_resultado,
                progress_bar,
                ventana,
                boton_transcribir,
                idioma_entrada_val,
                idioma_salida_val,
                diarizar_val,
                timestamps_val,
            ),
            kwargs={"spinner": spinner, "frame_progress": frame_progress, "on_file_done": on_file_done},
            daemon=True,
        ).start()



def procesar_audio(audio_file, idioma_entrada, translate, progress_bar, ventana, diarizar=False, status_cb=None, timestamps=False):
    """Transcribe a single file offline with faster-whisper.

    Whisper decodes the file itself (no ffmpeg conversion) and applies its own
    voice activity detection (no manual band-pass filtering or chunking).

    When ``diarizar`` is True, the speaker-labeled diarization path is used
    instead (whisperx + pyannote).
    """
    filename = os.path.basename(audio_file)

    def progress_cb(fraction):
        # Schedule on the main thread — Tkinter is not thread-safe.
        def _update():
            progress_bar.configure(value=fraction * 100)
            if status_cb:
                status_cb(fraction)
        ventana.after(0, _update)

    def should_continue():
        return config.transcripcion_activa

    if diarizar:
        transcripcion_final = diarizer.diarize_file(
            audio_file, language=idioma_entrada, progress_cb=progress_cb
        )
    elif timestamps:
        segments = config.transcription_service.transcribe_file_segments(
            audio_file,
            language=idioma_entrada,
            translate=translate,
            progress_cb=progress_cb,
            should_continue=should_continue,
        )
        parts = []
        for start, end, text in segments:
            parts.append(f"[{format_time(start)} - {format_time(end)}] {text}")
        transcripcion_final = "\n".join(parts)
    else:
        transcripcion_final = config.transcription_service.transcribe_file(
            audio_file,
            language=idioma_entrada,
            translate=translate,
            progress_cb=progress_cb,
            should_continue=should_continue,
        )

    if not transcripcion_final:
        transcripcion_final = "No se pudo transcribir ninguna parte del audio."

    palabras, inaudibles = contar_palabras_y_inaudibles(transcripcion_final)
    return {
        "filename": filename,
        "archivo": audio_file,
        "transcripcion": transcripcion_final,
        "inaudibles": inaudibles,
        "palabras": palabras,
    }


def iniciar_transcripcion(
    lista_archivos,
    text_area,
    archivo_procesando,
    lista_archivos_paths,
    transcripcion_resultado,
    progress_bar,
    ventana,
    boton_transcribir,
    idioma_entrada_val,
    idioma_salida_val,
    diarizar_val=False,
    timestamps_val=False,
    spinner=None,
    frame_progress=None,
    on_file_done=None,
):


    seleccion = lista_archivos.curselection()
    if not seleccion:
        show_warning(
            ventana,
            "Advertencia", "Por favor, seleccione al menos un archivo para transcribir.",
        )
        return

    diarizar = False

    archivos_seleccionados = [lista_archivos.get(i) for i in seleccion]
    total_archivos = len(archivos_seleccionados)

    idioma_entrada = idiomas[idioma_entrada_val]
    idioma_salida = idiomas[idioma_salida_val]

    # Offline translation: Whisper's translate task only targets English.
    # When the requested output is English (and the source is not), translate;
    # otherwise transcribe in the source language.
    translate = idioma_salida == "en" and idioma_entrada != "en"
    if (
        idioma_salida is not None
        and idioma_salida != "en"
        and idioma_salida != idioma_entrada
    ):
        show_info(
            ventana,
            "Traduccion offline",
            "La traduccion offline solo soporta ingles como destino. "
            "El audio se transcribira en su idioma original.",
        )

    config.transcripcion_activa = True
    config.transcripcion_en_curso = True
    ventana.after(0, lambda: boton_transcribir.configure(text="⏹   Detener"))

    # Pre-cargar el modelo si no está en memoria y avisar al usuario
    if config.transcription_service._model is None:
        archivo_procesando.set("Cargando modelo de transcripcion (solo la primera vez)...")
        try:
            config.transcription_service._get_model()
        except Exception as e:
            logger.error(f"Error al cargar el modelo: {str(e)}")

    for index, archivo in enumerate(archivos_seleccionados):
        if not config.transcripcion_activa:
            break

        audio_file = next(
            key for key, value in lista_archivos_paths.items() if value == archivo
        )

        # Consultar si el archivo ya fue transcrito
        record = config.repository.get(audio_file)
        if record:
            transcripcion_guardada = {
                "file_name": record.file_name,
                "duration": record.duration,
                "transcription": record.transcription,
                "summary": record.summary,
                "language": record.language
            }
        else:
            transcripcion_guardada = None
        if transcripcion_guardada:
            # Preguntar si desea cargar la transcripción existente o volver a procesar
            respuesta = _ask_on_main_thread(
                ventana,
                "Transcripción Guardada",
                f"El archivo '{archivo}' ya tiene una transcripción en el historial.\n\n"
                "¿Querés cargar la transcripción guardada?\n"
                "- Seleccioná SÍ para cargarla al instante.\n"
                "- Seleccioná NO para volver a transcribirla desde cero.\n"
                "- Seleccioná CANCELAR para abortar el proceso.",
            )
            if respuesta is True:
                nuevo_texto = (
                    f"Transcripcion de {archivo} (Historial guardado): \n"
                    f"{transcripcion_guardada['transcription']}\n\n"
                )
                ventana.after(0, lambda t=nuevo_texto: (
                    text_area.insert(tk.END, t),
                    text_area.see(tk.END),
                ))
                frac = ((index + 1) / total_archivos) * 100
                ventana.after(0, lambda v=frac: progress_bar.configure(value=v))
                continue
            elif respuesta is None:
                # Cancelar toda la cola
                break

        archivo_procesando.set(f"Procesando: {archivo} ({index + 1}/{total_archivos})")
        logger.info(f"Procesando archivo: {audio_file}")
        _t_start = time.time()
        _t_transcribe_start = None

        def status_cb(fraction):
            nonlocal _t_transcribe_start
            pct = int(fraction * 100)
            
            if fraction > 0.1001 and _t_transcribe_start is None:
                _t_transcribe_start = time.time()
                
            if _t_transcribe_start is not None and fraction > 0.1001:
                audio_fraction = min(1.0, (fraction - 0.1) / 0.8)
                if audio_fraction > 0.02:
                    elapsed_transcribing = time.time() - _t_transcribe_start
                    total_est = elapsed_transcribing / audio_fraction
                    remaining = total_est * (1.0 - audio_fraction)
                    if remaining > 120:
                        eta = f"~{int(remaining // 60)} min"
                    elif remaining > 60:
                        eta = "~1 min"
                    else:
                        eta = f"~{int(remaining)} seg"
                else:
                    eta = "calculando..."
            else:
                eta = "cargando..."
            archivo_procesando.set(f"Transcribiendo {pct}% ({eta}): {archivo} ({index + 1}/{total_archivos})")

        try:
            resultado = procesar_audio(
                audio_file, idioma_entrada, translate, progress_bar, ventana, diarizar, status_cb, timestamps=timestamps_val
            )

            if config.transcripcion_activa:
                nuevo_texto = (
                    f"Transcripcion de {archivo}: \n{resultado['transcripcion']}\n\n"
                    f"Palabras: {resultado['palabras']}\n\n"
                )
                ventana.after(0, lambda t=nuevo_texto: (
                    text_area.insert(tk.END, t),
                    text_area.see(tk.END),
                ))

                # Guardar en base de datos (no toca UI — seguro desde thread)
                from core.domain.entities import TranscriptionRecord
                saved_path = config.repository.save(
                    TranscriptionRecord(
                        file_path=audio_file,
                        file_name=resultado["filename"],
                        duration=archivo.split("(")[-1].replace(")", "").strip(),
                        transcription=resultado["transcripcion"],
                        summary="",
                        language=idioma_entrada or ""
                    )
                )

                if on_file_done:
                    ventana.after(0, lambda fp=saved_path: on_file_done(fp))

            frac = ((index + 1) / total_archivos) * 100
            ventana.after(0, lambda v=frac: progress_bar.configure(value=v))

        except diarizer.DiarizationError as e:
            _show_error_on_main_thread(
                ventana,
                "Diferenciar hablantes",
                report_error(f"Diarizacion {archivo}", e, str(e)),
            )
        except Exception as e:
            _show_error_on_main_thread(
                ventana,
                "Error",
                report_error(
                    f"Procesar {archivo}", e, f"No se pudo procesar el archivo {archivo}."
                ),
            )

        if not config.transcripcion_activa:
            break

    archivo_procesando.set("")

    if config.transcripcion_activa:
        _show_info_on_main_thread(
            ventana,
            "Informacion", f"Transcripcion completa para {total_archivos} archivo(s)."
        )

    def _reset_ui():
        boton_transcribir.configure(text="🎙️   Transcribir")
        progress_bar.configure(value=0)
        progress_bar.pack_forget()
        if spinner:
            spinner.stop()
            spinner.pack_forget()
        if frame_progress:
            frame_progress.pack_forget()

    ventana.after(0, _reset_ui)
    config.transcripcion_activa = False
    config.transcripcion_en_curso = False
